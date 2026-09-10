"""
Importa SOLO la hoja UNIDADES CHILE.
No lee otras hojas del Excel ni carpetas de otros proyectos.
Las fotos salen de public/cars/stock de ESTE repo, o de Drive filtrado por patente UC.
"""
from __future__ import annotations

import json
import re
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT_TS = ROOT / "src" / "data" / "cars.ts"
STOCK_DIR = ROOT / "public" / "cars" / "stock"
LOCAL_XLSX = Path(r"C:\Users\mathi\OneDrive\Escritorio\STOCK RG MOTORS_UNIDADES CHILE.xlsx")
SHEET_ID = "1BG2uR6APbXEMvVvRmdR-Nn0Vko6eobJ6Xam0XX41Ldc"
DRIVE_FOLDER_ID = "1etQDf-_InkLx8m4_AUMnc8xg2O_137St"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"

BRANDS = {
    "MITSUBISHI": "Mitsubishi",
    "TOYOTA": "Toyota",
    "PEUGEOT": "Peugeot",
    "NISSAN": "Nissan",
    "CHEVROLET": "Chevrolet",
    "FORD": "Ford",
    "VOLKSWAGEN": "Volkswagen",
    "VOLSWAGEN": "Volkswagen",
    "MAXUS": "Maxus",
    "MG": "MG",
    "SSANGYONG": "SsangYong",
    "HYUNDAI": "Hyundai",
    "RENAULT": "Renault",
    "MERCEDES": "Mercedes-Benz",
    "MERCEDEZ": "Mercedes-Benz",
    "SUBARU": "Subaru",
    "SUZUKI": "Suzuki",
    "KIA": "Kia",
    "JAC": "JAC",
    "CHERY": "Chery",
    "FIAT": "Fiat",
    "RAM": "RAM",
    "HINO": "Hino",
}


def clean_plate(value) -> str:
    return re.sub(r"[^A-Za-z0-9]", "", str(value or "")).upper()


def num(raw) -> int:
    if raw is None:
        return 0
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        n = int(round(float(raw)))
        if n > 100_000_000:
            n = round(n / 100)
        return n if n > 0 else 0
    s = str(raw)
    if re.search(r"falta|reservado|preparacion|taller|vendido|entregado|fotos", s, re.I):
        return 0
    digits = re.sub(r"[^0-9]", "", s.split("km")[0] if "km" in s.lower() else s)
    if not digits:
        return 0
    n = int(digits)
    if n > 100_000_000:
        n = round(n / 100)
    return n


def brand(raw) -> str:
    key = str(raw or "").strip().upper()
    return BRANDS.get(key, str(raw or "Otro").strip().title())


def body(model: str) -> str:
    m = model.upper()
    if re.search(r"HILUX|L200|KATANA|NAVARA|AMAROK|COLORADO|DMAX|T60|RANGER|SAVEIRO|PORTER", m):
        return "Pickup"
    if re.search(r"PARTNER|EXPERT|XZU|FURGON|FURGÓN", m):
        return "Furgón"
    if re.search(r"RAIZE|FORESTER|2008|CX|RAV|SUV", m):
        return "SUV"
    if re.search(r"SEDAN|SEDÁN", m):
        return "Sedán"
    return "Pickup"


def fuel(model: str) -> str:
    m = model.upper()
    if re.search(r"RAIZE|SAVEIRO|MSI", m):
        return "Bencina"
    return "Diésel"


def traction(model: str) -> str:
    m = model.upper()
    if "4X4" in m or "4WD" in m:
        return "4x4"
    return "4x2"


def cuota_desde(precio: int) -> int:
    if precio <= 0:
        return 0
    capital = precio * 0.8
    r = 0.0321
    n = 48
    return int(round(capital * (r * (1 + r) ** n) / ((1 + r) ** n - 1)))


def slugify(text: str) -> str:
    t = text.lower()
    t = re.sub(r"[^a-z0-9]+", "-", t).strip("-")
    return t


def load_rows():
    import openpyxl

    path = LOCAL_XLSX
    tmp = ROOT / ".cache" / "stock-live.xlsx"
    tmp.parent.mkdir(exist_ok=True)
    try:
        url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
        urllib.request.urlretrieve(url, tmp)
        if tmp.stat().st_size > 2000:
            path = tmp
    except Exception:
        path = LOCAL_XLSX

    wb = openpyxl.load_workbook(path, data_only=True)
    name = next((n for n in wb.sheetnames if n.strip().upper() == "UNIDADES CHILE"), None)
    if not name:
        raise SystemExit("No está la hoja UNIDADES CHILE")
    ws = wb[name]
    cars = []
    skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        plate = clean_plate(row[1])
        if len(plate) != 6 or not re.search(r"[A-Z]{2,4}[0-9]{2,4}", plate):
            skipped += 1
            continue
        blob = " ".join(str(x) for x in row if x is not None).upper()
        if any(k in blob for k in ("VENDIDO", "ENTREGADO", " VTA")):
            skipped += 1
            continue
        oferta = num(row[7])
        lista = num(row[6])
        precio = oferta or lista
        if precio <= 0:
            skipped += 1
            continue
        modelo = re.sub(r"\s+", " ", str(row[3] or "").strip())
        cars.append(
            {
                "plate": plate,
                "marca": brand(row[2]),
                "modelo": modelo,
                "version": modelo,
                "year": int(str(row[5]).split(".")[0]) if row[5] else 2022,
                "precio": precio,
                "mercado": lista if lista > precio else int(precio * 1.08),
                "km": num(row[8]),
                "color": str(row[4] or "").strip().title(),
            }
        )
    return cars, skipped, path == tmp


def fetch_text(url: str) -> str:
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=30) as res:
        return res.read().decode("utf-8", errors="ignore")


def extract_plate(text: str) -> str:
    compact = re.sub(r"[^A-Za-z0-9]", "", str(text or "")).upper()
    m = re.search(r"[A-Z]{4}\d{2}|[A-Z]{2}\d{4}|[A-Z]{3}\d{3}", compact)
    return m.group(0) if m else ""


def scrape_drive_folders(folder_id: str) -> dict[str, str]:
    html = fetch_text(f"https://drive.google.com/drive/folders/{folder_id}")
    found: dict[str, str] = {}
    patterns = [
        r'aria-label="([^"]+?)\s+(?:Shared folder|Google Drive Folder|folder)"[^>]*ssk=\'5:[^:]*:([a-zA-Z0-9_-]+)',
        r'\["([a-zA-Z0-9_-]{20,})",\["([^"\\]{3,80})"\]',
        r'\["([^"]{3,80})",\["([a-zA-Z0-9_-]{10,})"\]',
    ]
    for pat in patterns:
        for a, b in re.findall(pat, html, flags=re.I):
            name, fid = (b, a) if re.fullmatch(r"[a-zA-Z0-9_-]{20,}", a) and not re.fullmatch(r"[a-zA-Z0-9_-]{20,}", b) else (a, b)
            plate = extract_plate(name)
            if plate and plate not in found:
                found[plate] = fid.split("-")[0]
    # Fallback: any quoted plate-like folder title near a Drive id
    for name in re.findall(r'"([^"]{3,80})"', html):
        plate = extract_plate(name)
        if not plate or plate in found:
            continue
        window = html[max(0, html.find(name) - 120) : html.find(name) + 180]
        ids = re.findall(r"[-\w]{25,}", window)
        if ids:
            found[plate] = ids[0].lstrip("-")
    return found


def scrape_drive_photos(folder_id: str) -> list[str]:
    html = fetch_text(f"https://drive.google.com/drive/folders/{folder_id}")
    photos: list[tuple[str, str]] = []
    seen: set[str] = set()
    patterns = [
        r'aria-label="([^"]+?\.(?:jpg|jpeg|png|webp|heic))\s+Image\s+Shared"[^>]*ssk=\'5:[^:]*:([a-zA-Z0-9_-]+)',
        r'\["([^"]+\.(?:jpg|jpeg|png|webp|heic))",\["([a-zA-Z0-9_-]{10,})"\]',
    ]
    for pat in patterns:
        for name, fid in re.findall(pat, html, flags=re.I):
            file_id = fid.split("-")[0]
            if file_id in seen:
                continue
            seen.add(file_id)
            photos.append((name, file_id))
    photos.sort(key=lambda x: x[0].lower())
    return [fid for _, fid in photos[:12]]


def download_thumb(file_id: str, dest: Path) -> bool:
    url = f"https://drive.google.com/thumbnail?id={file_id}&sz=w1600"
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    try:
        with urllib.request.urlopen(req, timeout=30) as res:
            data = res.read()
        if len(data) < 800 or data[:15].lstrip().lower().startswith(b"<!doctype"):
            return False
        dest.write_bytes(data)
        return True
    except Exception:
        return False


def local_photos(plate: str) -> list[str]:
    dest = STOCK_DIR / plate.lower()
    if not dest.is_dir():
        return []
    return [
        f"/cars/stock/{plate.lower()}/{p.name}"
        for p in sorted(dest.iterdir())
        if p.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"}
    ]


def collect_photos(plate: str, drive_folders: dict[str, str]) -> list[str]:
    existing = local_photos(plate)
    if existing:
        return existing
    dest = STOCK_DIR / plate.lower()
    dest.mkdir(parents=True, exist_ok=True)
    copied = 0
    if plate in drive_folders:
        for i, file_id in enumerate(scrape_drive_photos(drive_folders[plate])):
            target = dest / f"{i:02d}.jpg"
            if download_thumb(file_id, target):
                copied += 1
    if copied:
        return local_photos(plate)
    return []


def ts_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace('"', '\\"')


def write_cars_ts(items: list[dict], with_photos: int):
    blocks = []
    for i, c in enumerate(items, 1):
        sid = slugify(f"{c['marca']}-{c['modelo']}-{c['year']}-{c['plate']}")
        imgs = c["imagenes"]
        img_lit = ", ".join(f'"{p}"' for p in imgs) if imgs else '"/cars/hero-l200.png"'
        trans = "Manual"
        if "AUTO" in c["modelo"].upper():
            trans = "Automático"
        blocks.append(
            f"""  {{
    id: "{sid}",
    unidad: "{c['plate'][:4]} {c['plate'][4:]}",
    marca: "{ts_escape(c['marca'])}",
    modelo: "{ts_escape(c['modelo'])}",
    version: "{ts_escape(c['version'])}",
    year: {c['year']},
    precio: {c['precio']},
    mercado: {c['mercado']},
    km: {c['km']},
    transmision: "{trans}",
    combustible: "{fuel(c['modelo'])}",
    traccion: "{traction(c['modelo'])}",
    duenos: 1,
    carroceria: "{body(c['modelo'])}",
    ciudad: "Puerto Montt",
    destacado: {str(i <= 3).lower()},
    certificado: true,
    cuota: {cuota_desde(c['precio'])},
    imagenes: [{img_lit}],
  }}"""
        )
    OUT_TS.write_text(
        """/** Stock de Unidades Chile Automotriz. */
export type Car = {
  id: string;
  unidad: string;
  marca: string;
  modelo: string;
  version: string;
  year: number;
  precio: number;
  mercado: number;
  km: number;
  transmision: "Automático" | "Manual";
  combustible: "Bencina" | "Diésel" | "Híbrido";
  traccion: "4x2" | "4x4";
  duenos: number;
  carroceria: "SUV" | "Sedán" | "Hatchback" | "Pickup" | "Crossover" | "Furgón";
  ciudad: string;
  destacado: boolean;
  certificado: boolean;
  cuota: number;
  imagenes: string[];
};

export const cars: Car[] = [
"""
        + ",\n".join(blocks)
        + """
];

export const marcas = [...new Set(cars.map((c) => c.marca))].sort();
export const modelos = [...new Set(cars.map((c) => c.modelo))].sort();
export const years = [...new Set(cars.map((c) => c.year))].sort((a, b) => b - a);
export const carrocerias = [...new Set(cars.map((c) => c.carroceria))];
export const combustibles = [...new Set(cars.map((c) => c.combustible))];

export function getCar(id: string) {
  return cars.find((c) => c.id === id);
}

export function featuredCars() {
  return cars.filter((c) => c.destacado);
}
""",
        encoding="utf-8",
    )


def main():
    items, skipped, live = load_rows()
    allowed = {c["plate"] for c in items}
    drive_folders: dict[str, str] = {}
    drive_ok = False
    try:
        scraped = scrape_drive_folders(DRIVE_FOLDER_ID)
        drive_folders = {plate: fid for plate, fid in scraped.items() if plate in allowed}
        drive_ok = bool(drive_folders)
    except Exception:
        drive_folders = {}
    matched_drive = sum(1 for c in items if c["plate"] in drive_folders)
    with_photos = 0
    for c in items:
        imgs = collect_photos(c["plate"], drive_folders)
        c["imagenes"] = imgs
        if imgs:
            with_photos += 1
    write_cars_ts(items, with_photos)
    summary = {
        "source": "google-live" if live else "local-excel",
        "active": len(items),
        "skipped": skipped,
        "driveFolders": len(drive_folders),
        "driveMatched": matched_drive,
        "driveOk": drive_ok,
        "withPhotos": with_photos,
        "withoutPhotos": len(items) - with_photos,
        "sheet": "UNIDADES CHILE",
    }
    (ROOT / ".cache" / "import-stock-summary.json").write_text(
        json.dumps(summary, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary))


if __name__ == "__main__":
    main()
